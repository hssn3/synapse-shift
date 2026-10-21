"""استخراج رستری کامل کد ۵۵ و ۹۹ از فایل‌های Word."""
import sys, json, re
from docx import Document

ROLES = [
    ('1', 'پزشک طب اورژانس'),
    ('2', 'سوپروایزر'),
    ('3', 'تکنسین بیهوشی'),
    ('4', 'پرستار'),
    ('5', 'پرستار'),
    ('6', 'پرستار'),
]
SHIFTS = [('M', 'صبح'), ('E', 'عصر'), ('N', 'شب')]


def cell(c):
    return ' '.join(p.text.strip() for p in c.paragraphs if p.text.strip()).strip()


def parse(path):
    doc = Document(path)
    out = {}
    for tbl in doc.tables:
        rows = tbl.rows
        if len(rows) < 20:
            continue
        header = [cell(c) for c in rows[1].cells]
        dates = [h for h in header if '/' in h]
        if not dates:
            continue
        # ایندکس ستون‌های روز در هر ردیف
        for ri in range(2, len(rows)):
            cells = [cell(c) for c in rows[ri].cells]
            shift_label = cells[0]
            role_label = cells[1]
            if shift_label not in ('صبح', 'عصر', 'شب'):
                continue
            shift_key = {'صبح': 'M', 'عصر': 'E', 'شب': 'N'}[shift_label]
            role_no = None
            # آخرین رقم ۱..۶ در برچسب نقش: «… 1»، «(A-3)»، «B-4» و مانند آن
            digits = re.findall(r'[1-6]', role_label)
            if digits:
                role_no = digits[-1]
            if role_no is None:
                continue
            vals = cells[2:2 + len(dates)]
            for d, v in zip(dates, vals):
                day = int(d.split('/')[0])
                out.setdefault(shift_key, {}).setdefault(role_no, {})[day] = v.strip()
    return out


if __name__ == '__main__':
    result = {}
    for path in sys.argv[1:]:
        key = '55' if '55' in path else ('99' if '99' in path else None)
        if not key:
            continue
        result[key] = parse(path)
    with open('build/code-rosters.json', 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=1)
    for k, v in result.items():
        days = sorted({d for sh in v.values() for r in sh.values() for d in r})
        print('code', k, 'shifts', sorted(v), 'days', len(days), days[:3], '...', days[-2:])
        for sh in ('M', 'E', 'N'):
            for role in ('1', '2', '3', '4', '5', '6'):
                vals = v.get(sh, {}).get(role, {})
                uniq = sorted(set(vals.values()))
                print('   ', sh, role, 'cells', len(vals), 'uniq', len(uniq), uniq[:4])
