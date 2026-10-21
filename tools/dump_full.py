import sys, os, io, difflib
from docx import Document

def cell_text(cell):
    return ' '.join(p.text.strip() for p in cell.paragraphs if p.text.strip())

def table_text(t, depth=0):
    out = []
    for r in t.rows:
        seen = set()
        vals = []
        for c in r.cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            vals.append(cell_text(c))
        out.append(' | '.join(vals))
    return out

def doc_dump(path):
    d = Document(path)
    lines = []
    for s in d.sections:
        for t in s.header.tables:
            lines += ['HDR: ' + x for x in table_text(t)]
        for p in s.header.paragraphs:
            if p.text.strip():
                lines.append('HDR-P: ' + p.text.strip())
    for t in d.tables:
        lines += table_text(t)
    for p in d.paragraphs:
        if p.text.strip():
            lines.append('P: ' + p.text.strip())
    return lines

if __name__ == '__main__':
    paths = sys.argv[1:]
    dumps = {}
    for p in paths:
        name = os.path.basename(p)
        dumps[name] = doc_dump(p)
        print('=' * 78)
        print(name, '-> %d lines' % len(dumps[name]))
        for l in dumps[name][:25]:
            print('  ', l[:180])
        print('   ...')
        for l in dumps[name][-10:]:
            print('  ', l[:180])
