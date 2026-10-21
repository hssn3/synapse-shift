"""Dump structure of DOCX files: sections, styles, paragraphs, tables, images."""
import sys, json, os
from docx import Document
from docx.shared import Emu

def walk_blocks(doc):
    """Yield body blocks in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            yield ('p', Paragraph(child, doc))
        elif tag == 'tbl':
            yield ('tbl', Table(child, doc))

def para_info(p):
    text = p.text.strip()
    style = p.style.name if p.style is not None else ''
    return {
        'style': style,
        'align': str(p.alignment).split('.')[-1].split(' ')[0] if p.alignment is not None else '',
        'text': text,
    }

def table_info(t):
    rows = []
    header = []
    if len(t.rows):
        header = [c.text.strip().replace('\n', ' / ') for c in t.rows[0].cells]
    ncols = max((len(r.cells) for r in t.rows), default=0)
    for r in t.rows[:6]:
        rows.append([c.text.strip().replace('\n', ' / ')[:60] for c in r.cells])
    return {'rows': len(t.rows), 'cols': ncols, 'header': header, 'sample': rows}

def inspect(path):
    doc = Document(path)
    out = {'file': os.path.basename(path), 'size': os.path.getsize(path)}
    secs = []
    for s in doc.sections:
        secs.append({
            'page': f"{round(s.page_width.cm,1)}x{round(s.page_height.cm,1)} cm",
            'orientation': 'landscape' if s.page_width > s.page_height else 'portrait',
            'margins_cm': [round(s.top_margin.cm,1), round(s.right_margin.cm,1),
                           round(s.bottom_margin.cm,1), round(s.left_margin.cm,1)],
            'header': ' | '.join(p.text.strip() for p in s.header.paragraphs if p.text.strip())[:120],
            'footer': ' | '.join(p.text.strip() for p in s.footer.paragraphs if p.text.strip())[:120],
        })
    out['sections'] = secs
    out['core'] = {
        'title': doc.core_properties.title or '',
        'author': doc.core_properties.author or '',
        'created': str(doc.core_properties.created or ''),
    }

    paras, tables, blocks = [], [], []
    for kind, obj in walk_blocks(doc):
        if kind == 'p':
            info = para_info(obj)
            if info['text']:
                paras.append(info)
                blocks.append({'k': 'p', **info})
        else:
            info = table_info(obj)
            tables.append(info)
            blocks.append({'k': 'tbl', 'rows': info['rows'], 'cols': info['cols'],
                           'header': info['header']})
    out['counts'] = {
        'paragraphs_with_text': len(paras),
        'tables': len(tables),
        'inline_images': len(doc.inline_shapes),
        'blocks': len(blocks),
    }
    styles = {}
    for p in paras:
        styles[p['style']] = styles.get(p['style'], 0) + 1
    out['style_histogram'] = dict(sorted(styles.items(), key=lambda x: -x[1]))
    out['headings'] = [b['text'] for b in blocks if b['k'] == 'p' and ('Head' in b['style'] or 'Title' in b['style'] or 'عنوان' in b['style'])]
    out['all_text_first_60'] = [b['text'][:100] for b in blocks if b['k'] == 'p'][:60]
    out['tables_detail'] = tables
    out['blocks'] = blocks[:120]
    return out

if __name__ == '__main__':
    results = [inspect(p) for p in sys.argv[1:]]
    print(json.dumps(results, ensure_ascii=False, indent=1)[:200000])
